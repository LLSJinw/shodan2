"""Deterministic Excel and Word report generation for PQC Recon V2."""

from __future__ import annotations

from collections import Counter
import datetime as dt
from io import BytesIO
import re
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl.styles import Alignment, Font, PatternFill


INK = "192C36"
TEAL = "087F73"
PALE = "E7F3F1"
MUTED = "5D6C74"
RED = "B74D46"
PRIORITY_STYLES = {
    "Urgent": ("F4D8D6", "8B2F2A"),
    "High": ("F8E5CF", "8A5314"),
    "Medium": ("F5EFD9", "6F5A12"),
    "Review": ("E8EEF0", "455B65"),
    "Low": ("E8F3EF", "27685F"),
    "Info": ("EEF2F3", "52636C"),
}
PRIORITY_ORDER = {"Urgent": 0, "High": 1, "Medium": 2, "Review": 3}
AI_INPUT_COLUMNS = [
    "Section",
    "Order",
    "Metric",
    "Value",
    "Priority",
    "Asset",
    "CVE",
    "CVSS",
    "EPSS",
    "KEV",
    "Detail",
    "Source Sheet",
    "Source Rule",
]


def _frame(rows: list[dict[str, Any]], columns: list[str] | None = None) -> pd.DataFrame:
    frame = pd.DataFrame(rows)
    if columns:
        for column in columns:
            if column not in frame:
                frame[column] = ""
        frame = frame[columns]
    return frame


def _as_float(value: Any, default: float = -1.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _cve_year(value: Any) -> int:
    match = re.match(r"^CVE-(\d{4})-", str(value or ""), flags=re.IGNORECASE)
    return int(match.group(1)) if match else 0


def _split_csv(value: Any) -> list[str]:
    return [item.strip() for item in str(value or "").split(",") if item.strip()]


def _is_true(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"true", "yes", "1"}


def _parse_date(value: Any) -> dt.date | None:
    if value in (None, ""):
        return None
    try:
        return dt.date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _ai_input_rows(result: dict[str, Any]) -> list[dict[str, Any]]:
    """Build a canonical, model-readable manifest without generative interpretation."""
    findings = [row for row in result.get("findings", []) if row.get("Finding ID")]
    assets = [row for row in result.get("asset_rows", []) if row.get("IP")]
    cves = [row for row in result.get("cve_rows", []) if row.get("CVE ID")]
    tls_rows = list(result.get("tls_rows", []))
    source_health = [row for row in result.get("source_health", []) if any(value not in (None, "") for value in row.values())]
    rejected = [row for row in result.get("rejected_targets", []) if any(value not in (None, "") for value in row.values())]

    rows: list[dict[str, Any]] = []

    def add(
        section: str,
        order: int,
        metric: str,
        value: Any = "",
        *,
        priority: str = "",
        asset: str = "",
        cve: str = "",
        cvss: Any = "",
        epss: Any = "",
        kev: str = "",
        detail: str = "",
        source_sheet: str = "",
        source_rule: str = "",
    ) -> None:
        rows.append({
            "Section": section,
            "Order": order,
            "Metric": metric,
            "Value": value,
            "Priority": priority,
            "Asset": asset,
            "CVE": cve,
            "CVSS": cvss,
            "EPSS": epss,
            "KEV": kev,
            "Detail": detail,
            "Source Sheet": source_sheet,
            "Source Rule": source_rule,
        })

    add(
        "USAGE_RULE",
        1,
        "Canonical values",
        "Use exactly",
        detail="Copy counts and ranked CVEs from AI_INPUT. Do not recalculate or add report rows to source totals.",
        source_sheet="AI_INPUT",
        source_rule="Application-generated deterministic manifest",
    )
    add(
        "USAGE_RULE",
        2,
        "Evidence boundary",
        "External observations only",
        detail="Do not claim confirmed ownership, vulnerability applicability, exploitability, control gaps, business risk, or product need.",
        source_sheet="Methodology",
        source_rule="Customer and authenticated validation required",
    )
    add(
        "USAGE_RULE",
        3,
        "Raw appendix",
        "All non-AI_INPUT sheets",
        detail="The original workbook sheets remain Appendix A. AI_INPUT is a canonical index, not a replacement for raw evidence.",
        source_sheet="Workbook",
        source_rule="Do not reproduce all raw rows in the main report",
    )

    hostnames = [hostname for row in assets for hostname in _split_csv(row.get("Hostnames"))]
    service_pairs: set[tuple[str, int]] = set()
    for row in assets:
        ports = row.get("Ports")
        if not isinstance(ports, list):
            ports = _split_csv(row.get("Open TCP ports"))
        for port in ports:
            try:
                service_pairs.add((str(row.get("IP", "")), int(port)))
            except (TypeError, ValueError):
                continue

    scope_counts = [
        ("Public IPs", len({str(row.get("IP")) for row in assets})),
        ("Hostname records", len(hostnames)),
        ("Unique hostnames", len(set(hostnames))),
        ("Open-service observations", len(service_pairs)),
        ("Rejected targets", len(rejected)),
    ]
    for order, (metric, value) in enumerate(scope_counts, start=1):
        source = "Rejected_Targets" if metric == "Rejected targets" else "Assets"
        add("SCOPE_COUNT", order, metric, value, source_sheet=source, source_rule="Counted from nonblank raw records")

    finding_priorities = Counter(str(row.get("Priority") or "Review") for row in findings)
    add("OVERALL_FINDINGS", 1, "Total findings", len(findings), source_sheet="Priority_Findings", source_rule="Nonblank Finding ID rows")
    for order, priority in enumerate(PRIORITY_ORDER, start=2):
        add(
            "OVERALL_FINDINGS",
            order,
            f"{priority} findings",
            finding_priorities.get(priority, 0),
            priority=priority,
            source_sheet="Priority_Findings",
            source_rule="Group nonblank Finding ID rows by Priority",
        )

    cve_priorities = Counter(str(row.get("Technical Priority") or "Review") for row in cves)
    unique_cves = {str(row.get("CVE ID")) for row in cves}
    add("CVE_COUNTS", 1, "CVE associations", len(cves), source_sheet="Vulnerabilities", source_rule="Nonblank IP + CVE ID rows")
    add("CVE_COUNTS", 2, "Unique CVE IDs", len(unique_cves), source_sheet="Vulnerabilities", source_rule="Unique nonblank CVE ID values")
    add("CVE_COUNTS", 3, "CISA KEV Yes", sum(1 for row in cves if str(row.get("CISA KEV")) == "Yes"), source_sheet="Vulnerabilities", source_rule="Exact Yes values")
    add("CVE_COUNTS", 4, "Known ransomware use Yes", sum(1 for row in cves if str(row.get("Known ransomware use")) == "Yes"), source_sheet="Vulnerabilities", source_rule="Exact Yes values")
    add("CVE_COUNTS", 5, "Known ransomware use Unknown", sum(1 for row in cves if str(row.get("Known ransomware use")) == "Unknown"), source_sheet="Vulnerabilities", source_rule="Exact Unknown values")
    for order, priority in enumerate(PRIORITY_ORDER, start=6):
        add(
            "CVE_PRIORITY",
            order,
            f"{priority} CVE associations",
            cve_priorities.get(priority, 0),
            priority=priority,
            source_sheet="Vulnerabilities",
            source_rule="Group nonblank CVE rows by Technical Priority",
        )

    completed_date = _parse_date(result.get("meta", {}).get("completed_at")) or dt.datetime.now(dt.timezone.utc).date()
    reachable = [row for row in tls_rows if _is_true(row.get("reachable"))]
    failed = [row for row in tls_rows if not _is_true(row.get("reachable"))]
    certificate_status = Counter()
    for row in tls_rows:
        if not _is_true(row.get("reachable")):
            certificate_status["Unknown"] += 1
            continue
        not_before = _parse_date(row.get("not_before"))
        days_to_expiry = _as_float(row.get("days_to_expiry"), default=float("nan"))
        if not_before and not_before > completed_date:
            certificate_status["Not yet valid"] += 1
        elif days_to_expiry != days_to_expiry:
            certificate_status["Unknown"] += 1
        elif days_to_expiry < 0:
            certificate_status["Expired"] += 1
        else:
            certificate_status["Valid"] += 1

    tls_counts = [
        ("TLS endpoint combinations", len(tls_rows)),
        ("TLS successful", len(reachable)),
        ("TLS failed", len(failed)),
        ("Certificates valid", certificate_status.get("Valid", 0)),
        ("Certificates expired", certificate_status.get("Expired", 0)),
        ("Certificates not yet valid", certificate_status.get("Not yet valid", 0)),
        ("Certificate status unknown", certificate_status.get("Unknown", 0)),
        ("RSA certificate observations", sum(1 for row in reachable if str(row.get("public_key_algorithm") or "").upper().startswith("RSA"))),
        ("Quantum-vulnerable certificate observations", sum(1 for row in reachable if _is_true(row.get("quantum_vulnerable")))),
    ]
    for order, (metric, value) in enumerate(tls_counts, start=1):
        add("TLS_COUNTS", order, metric, value, source_sheet="TLS_PQC", source_rule="Endpoint rows; failed handshakes have unknown certificate status")

    kex_counts = Counter(str(row.get("key_exchange_pqc") or "UNKNOWN").upper() for row in reachable)
    kex_rows = [
        ("PQC-protected KEX", kex_counts.get("PQC-PROTECTED", 0)),
        ("Classical-only KEX", kex_counts.get("CLASSICAL-ONLY", 0)),
        ("Unknown KEX", kex_counts.get("UNKNOWN", 0)),
        ("Failed KEX tests", len(failed)),
    ]
    for order, (metric, value) in enumerate(kex_rows, start=1):
        add("KEX_COUNTS", order, metric, value, source_sheet="TLS_PQC", source_rule="Successful rows grouped by key_exchange_pqc; failed rows separate")

    source_counts = [
        ("Priority_Findings rows", len(findings), "Priority_Findings"),
        ("Assets rows", len(assets), "Assets"),
        ("Vulnerabilities rows", len(cves), "Vulnerabilities"),
        ("TLS_PQC rows", len(tls_rows), "TLS_PQC"),
        ("Source_Health rows", len(source_health), "Source_Health"),
        ("Rejected_Targets rows", len(rejected), "Rejected_Targets"),
    ]
    for order, (metric, value, source) in enumerate(source_counts, start=1):
        add("SOURCE_COVERAGE", order, metric, value, source_sheet=source, source_rule="Nonblank data rows")

    add(
        "RECONCILIATION",
        1,
        "Overall priority sum equals total findings",
        "PASS" if sum(finding_priorities.get(priority, 0) for priority in PRIORITY_ORDER) == len(findings) else "FAIL",
        source_sheet="Priority_Findings",
        source_rule="Urgent + High + Medium + Review = Total findings",
    )
    add(
        "RECONCILIATION",
        2,
        "CVE priority sum equals CVE associations",
        "PASS" if sum(cve_priorities.get(priority, 0) for priority in PRIORITY_ORDER) == len(cves) else "FAIL",
        source_sheet="Vulnerabilities",
        source_rule="Urgent + High + Medium + Review = CVE associations",
    )
    add(
        "RECONCILIATION",
        3,
        "TLS status sum equals endpoint combinations",
        "PASS" if sum(certificate_status.values()) == len(tls_rows) else "FAIL",
        source_sheet="TLS_PQC",
        source_rule="Valid + Expired + Not yet valid + Unknown = TLS endpoint combinations",
    )
    add(
        "RECONCILIATION",
        4,
        "KEX status sum equals endpoint combinations",
        "PASS" if sum(kex_counts.values()) + len(failed) == len(tls_rows) else "FAIL",
        source_sheet="TLS_PQC",
        source_rule="PQC-protected + Classical-only + Unknown + Failed = TLS endpoint combinations",
    )

    ranked_cves = sorted(
        cves,
        key=lambda row: (
            0 if str(row.get("CISA KEV")) == "Yes" else 1,
            0 if str(row.get("Known ransomware use")) == "Yes" else 1,
            PRIORITY_ORDER.get(str(row.get("Technical Priority") or "Review"), 9),
            -_as_float(row.get("CVSS")),
            -_as_float(row.get("EPSS")),
            -_cve_year(row.get("CVE ID")),
            str(row.get("IP") or ""),
            str(row.get("CVE ID") or ""),
        ),
    )
    for rank, row in enumerate(ranked_cves[:10], start=1):
        add(
            "TOP_CVE",
            rank,
            "Priority CVE association",
            rank,
            priority=str(row.get("Technical Priority") or "Review"),
            asset=str(row.get("IP") or ""),
            cve=str(row.get("CVE ID") or ""),
            cvss=row.get("CVSS", ""),
            epss=row.get("EPSS", ""),
            kev=str(row.get("CISA KEV") or "Unknown"),
            detail=str(row.get("Title") or "Description unavailable"),
            source_sheet="Vulnerabilities",
            source_rule="KEV, ransomware, priority, CVSS, EPSS, CVE year, asset, CVE ID",
        )

    for order, (ip, port) in enumerate(sorted(service_pairs), start=1):
        add(
            "OPEN_SERVICE",
            order,
            "Observed IP-port pair",
            port,
            asset=ip,
            detail="Externally reported service observation; product, purpose, ownership, and control context unconfirmed.",
            source_sheet="Assets",
            source_rule="Unique IP + port pair",
        )

    return rows


def build_excel(result: dict[str, Any]) -> BytesIO:
    meta = result.get("meta", {})
    summary = result.get("summary", {})
    overview_rows = [{"Field": key.replace("_", " ").title(), "Value": value} for key, value in {**meta, **summary}.items()]
    asset_rows = [{key: value for key, value in row.items() if key != "Ports"} for row in result.get("asset_rows", [])]
    sheets = {
        "AI_INPUT": _frame(_ai_input_rows(result), AI_INPUT_COLUMNS),
        "Overview": _frame(overview_rows),
        "Priority_Findings": _frame(result.get("findings", [])),
        "Assets": _frame(asset_rows),
        "Vulnerabilities": _frame(result.get("cve_rows", [])),
        "TLS_PQC": _frame(result.get("tls_rows", [])),
        "Source_Health": _frame(result.get("source_health", [])),
        "Rejected_Targets": _frame(result.get("rejected_targets", [])),
    }

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, index=False, sheet_name=sheet_name[:31])
            sheet = writer.book[sheet_name[:31]]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor=INK)
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center")
            for column in sheet.columns:
                values = [str(cell.value or "") for cell in column[:100]]
                width = min(max(max(map(len, values), default=8) + 2, 10), 48)
                sheet.column_dimensions[column[0].column_letter].width = width
                for cell in column:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
    output.seek(0)
    return output


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _shade_run(run, fill: str) -> None:
    properties = run._r.get_or_add_rPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_table_borders(table, color: str = "D9D9D9") -> None:
    properties = table._tbl.tblPr
    existing = properties.find(qn("w:tblBorders"))
    if existing is not None:
        properties.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)
        borders.append(border)
    properties.append(borders)


def _set_cell_text(cell, value: Any, *, bold: bool = False, color: str = INK, size: int = 8):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(value if value not in (None, "") else "-"))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    return run


def _add_table(
    document: Document,
    rows: list[dict[str, Any]],
    columns: list[str],
    limit: int = 30,
    styled_columns: dict[str, dict[str, tuple[str, str]]] | None = None,
) -> None:
    if not rows:
        document.add_paragraph("No observations recorded.", style="Intense Quote")
        return
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    _set_table_borders(table)
    for index, column in enumerate(columns):
        _set_cell_text(table.rows[0].cells[index], column, bold=True, color="FFFFFF")
        _shade(table.rows[0].cells[index], INK)
    for row in rows[:limit]:
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            value = row.get(column, "")
            run = _set_cell_text(cells[index], value)
            palette = (styled_columns or {}).get(column, {}).get(str(value))
            if palette:
                fill, color = palette
                _shade(cells[index], fill)
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(color)
            elif len(table.rows) % 2 == 0:
                _shade(cells[index], "F4F7F8")
    if len(rows) > limit:
        note = document.add_paragraph(f"Showing {limit} of {len(rows)} records. See the Excel evidence package for all rows.")
        note.style = "Caption"


def _heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(5)


def _remove_style_border(style) -> None:
    paragraph_properties = style.element.get_or_add_pPr()
    border = paragraph_properties.find(qn("w:pBdr"))
    if border is not None:
        paragraph_properties.remove(border)


def build_docx(result: dict[str, Any]) -> BytesIO:
    meta = result.get("meta", {})
    summary = result.get("summary", {})
    findings = result.get("findings", [])
    assets = result.get("asset_rows", [])
    cves = result.get("cve_rows", [])
    tls_rows = result.get("tls_rows", [])

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(26)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    _remove_style_border(styles["Title"])
    styles["Subtitle"].font.name = "Aptos"
    styles["Subtitle"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Subtitle"].font.italic = False
    _remove_style_border(styles["Subtitle"])
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph(style="Title")
    title.add_run("External Attack Surface and PQC Readiness")
    subtitle = document.add_paragraph("Presales Reconnaissance and Customer Validation Report")
    subtitle.style = "Subtitle"

    identity_rows = (
        ("Customer / account", meta.get("customer", "Internal or sanitized assessment")),
        ("Engagement", meta.get("engagement", "External reconnaissance")),
        ("Assessment profile", meta.get("profile", "Standard presales")),
        ("Assessment time", meta.get("completed_at", "")),
        ("Scope", meta.get("scope", "Authorized targets supplied for this run")),
    )
    identity = document.add_table(rows=len(identity_rows), cols=2)
    identity.style = "Table Grid"
    _set_table_borders(identity)
    for row, values in zip(identity.rows, identity_rows):
        _set_cell_text(row.cells[0], values[0], bold=True, color=TEAL, size=9)
        _set_cell_text(row.cells[1], values[1], size=9)
        _shade(row.cells[0], PALE)

    notice = document.add_paragraph()
    notice.paragraph_format.space_before = Pt(8)
    notice.paragraph_format.space_after = Pt(8)
    notice_lead = notice.add_run("Interpretation boundary: ")
    notice_lead.bold = True
    notice.add_run(
        "External telemetry is an observation at a point in time. CVE associations, ownership, service exposure, "
        "and business impact require customer validation before remediation or solution decisions."
    )

    _heading(document, "Executive snapshot")
    metrics = [
        ("Public IPs observed", summary.get("public_ips", 0)),
        ("Reported open services", summary.get("open_ports", 0)),
        ("CVE associations", summary.get("cve_associations", 0)),
        ("CISA KEV matches", summary.get("kev_matches", 0)),
        ("TLS endpoints assessed", summary.get("tls_endpoints", 0)),
        ("Classical-only key exchange", summary.get("classical_only_kex", 0)),
        ("PQC-protected key exchange", summary.get("pqc_protected_kex", 0)),
        ("Unknown key exchange", summary.get("unknown_kex", 0)),
    ]
    metric_table = document.add_table(rows=2, cols=4)
    metric_table.style = "Table Grid"
    _set_table_borders(metric_table)
    for index, (label, value) in enumerate(metrics):
        cell = metric_table.rows[index // 4].cells[index % 4]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        score_run = paragraph.add_run(f"{value}\n")
        score_run.bold = True
        score_run.font.size = Pt(18)
        score_run.font.color.rgb = RGBColor.from_string(TEAL)
        label_run = paragraph.add_run(label)
        label_run.font.size = Pt(7)
        label_run.font.color.rgb = RGBColor.from_string(MUTED)
        _shade(cell, PALE)

    priority_counts = Counter(str(finding.get("Priority", "Review")) for finding in findings)
    priority_table = document.add_table(rows=1, cols=4)
    priority_table.style = "Table Grid"
    _set_table_borders(priority_table)
    for index, level in enumerate(("Urgent", "High", "Medium", "Review")):
        fill, color = PRIORITY_STYLES[level]
        cell = priority_table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        count_run = paragraph.add_run(f"{priority_counts.get(level, 0)}\n")
        count_run.bold = True
        count_run.font.size = Pt(16)
        count_run.font.color.rgb = RGBColor.from_string(color)
        label_run = paragraph.add_run(level)
        label_run.bold = True
        label_run.font.size = Pt(8)
        label_run.font.color.rgb = RGBColor.from_string(color)
        _shade(cell, fill)

    _heading(document, "Prioritized findings")
    if not findings:
        document.add_paragraph("No prioritized findings were generated from the available observations.")
    for finding in findings[:20]:
        finding_title = document.add_paragraph()
        finding_title.paragraph_format.space_before = Pt(7)
        finding_title.paragraph_format.space_after = Pt(3)
        identifier = finding_title.add_run(f"{finding.get('Finding ID', '')}  ")
        identifier.bold = True
        identifier.font.size = Pt(11)
        priority = str(finding.get("Priority", "Review"))
        fill, color = PRIORITY_STYLES.get(priority, PRIORITY_STYLES["Review"])
        priority_run = finding_title.add_run(f" {priority.upper()} ")
        priority_run.bold = True
        priority_run.font.size = Pt(9)
        priority_run.font.color.rgb = RGBColor.from_string(color)
        _shade_run(priority_run, fill)
        domain = finding_title.add_run(f"  {finding.get('Domain', '')}")
        domain.bold = True
        domain.font.size = Pt(11)
        document.add_paragraph(f"Asset: {finding.get('Asset', '-')}")
        document.add_paragraph(str(finding.get("Observation", "")))
        for label, key in (
            ("Evidence", "Evidence"),
            ("Confidence", "Confidence"),
            ("Why it matters", "Why it matters"),
            ("Customer validation", "Customer validation"),
            ("Suggested next step", "Suggested next step"),
        ):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            label_run = paragraph.add_run(f"{label}: ")
            label_run.bold = True
            paragraph.add_run(str(finding.get(key, "-") or "-"))
    if len(findings) > 20:
        document.add_paragraph(f"Showing 20 of {len(findings)} findings. See the Excel evidence package for all records.")

    _heading(document, "External asset inventory")
    asset_report_rows = [{**row, "Open TCP ports": row.get("Open TCP ports", "")} for row in assets]
    _add_table(document, asset_report_rows, ["IP", "Hostnames", "Discovery sources", "Open TCP ports", "InternetDB status"], limit=40)

    _heading(document, "Vulnerability observations")
    _add_table(
        document,
        cves,
        ["IP", "CVE ID", "CVSS", "CISA KEV", "EPSS", "Technical Priority", "Title"],
        limit=35,
        styled_columns={"Technical Priority": PRIORITY_STYLES},
    )

    _heading(document, "TLS and post-quantum readiness")
    tls_report_rows = []
    for row in tls_rows:
        tls_report_rows.append({
            "Endpoint": f"{row.get('hostname') or row.get('ip')}:{row.get('port', '')}",
            "TLS": row.get("tls_version", ""),
            "Certificate expiry": row.get("not_after", ""),
            "Public key": f"{row.get('public_key_algorithm', '')} {row.get('public_key_size') or ''}".strip(),
            "KEX verdict": row.get("key_exchange_pqc", ""),
            "Risk": row.get("risk", ""),
            "Findings": row.get("findings", ""),
        })
    _add_table(
        document,
        tls_report_rows,
        ["Endpoint", "TLS", "Certificate expiry", "Public key", "KEX verdict", "Risk", "Findings"],
        limit=35,
        styled_columns={"Risk": PRIORITY_STYLES},
    )

    _heading(document, "Interpretation and limitations")
    limitations = (
        "Shodan InternetDB and DNS sources may be incomplete or stale; absence of a reported CVE is not evidence of security.\n"
        "CVE-to-asset associations require authenticated product and version validation.\n"
        "A normal TLS handshake does not test every protocol version or cipher supported by a service.\n"
        "CA/Browser Forum validity limits apply to publicly trusted subscriber certificates; trust applicability requires validation.\n"
        "NIST IR 8547 transition dates are sourced from an Initial Public Draft and must be labeled accordingly.\n"
        "PQC urgency depends on data sensitivity, confidentiality lifetime, architecture, and migration dependencies."
    )
    for line in limitations.splitlines():
        document.add_paragraph(line, style="List Bullet")

    _heading(document, "Data-source health")
    _add_table(document, result.get("source_health", []), ["Source", "Status", "Detail"], limit=20)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
